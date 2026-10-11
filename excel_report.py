from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def create_excel_report(total_sales, sales_by_product, output_file="sales_report.xlsx"):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sales Report"
    sheet.append(["Product", "Total Sales"])
    for product, sales in sales_by_product.items():
        sheet.append([product, sales])
    sheet.append(["Total Sales", total_sales])

    for cell in sheet[1]:
         cell.font = Font(bold=True, color="FFFFFF")
         cell.fill = PatternFill("solid", fgColor="1F4E78")
         cell.alignment = Alignment(horizontal="center")

    for cell in sheet[sheet.max_row]:
         cell.font = Font(bold=True)

    for column in sheet.columns:
         max_lentgh = max(len(str(cell.value or ""))for cell in column)
         letter = get_column_letter(column[0].column)
         sheet.column_dimensions[letter].width = max_lentgh + 4
    for row in range(2, sheet.max_row + 1):
           sheet.cell(row=row, column=2).number_format = '#,##0'
    workbook.save(output_file)